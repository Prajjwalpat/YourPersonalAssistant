"""Offline regression checks for the first assessment improvement batch."""
import io
import sys
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from docx import Document as WordDocument
from langchain_core.documents import Document
from langchain_core.retrievers import BaseRetriever
from streamlit.testing.v1 import AppTest

from conversation import get_sources
from document_utils import ingest_files, parse_uploaded_file
from text_processing import get_chunks, get_hybrid_retriever


class Upload(io.BytesIO):
    def __init__(self, name, content, mime="application/octet-stream"):
        super().__init__(content)
        self.name, self.type = name, mime


class StaticRetriever(BaseRetriever):
    documents: list

    def _get_relevant_documents(self, query, *, run_manager):
        return self.documents


class ParserTests(unittest.TestCase):
    def test_markdown_and_code_use_extension_and_can_be_read_again(self):
        for name in ("notes.md", "notes.markdown", "main.py", "query.sql", "widget.tsx"):
            upload = Upload(name, b"\xef\xbb\xbf# Heading\nCODE-42 = 7")
            upload.read()  # Simulate an already-consumed upload stream.
            first = parse_uploaded_file(upload)
            self.assertEqual(first[0].metadata["source"], name)
            self.assertEqual(first[0].page_content, "# Heading\nCODE-42 = 7")
            self.assertEqual(parse_uploaded_file(upload), first)

    def test_pdf_page_numbers_survive_chunking_and_boundaries(self):
        pages = [SimpleNamespace(extract_text=lambda: "First page. " * 90),
                 SimpleNamespace(extract_text=lambda: "Second page."),
                 SimpleNamespace(extract_text=lambda: None)]
        with patch("document_utils.PdfReader", return_value=SimpleNamespace(pages=pages)):
            documents = parse_uploaded_file(Upload("report.pdf", b"fake PDF"))
        chunks = get_chunks(documents)
        self.assertEqual({c.metadata["page"] for c in chunks}, {1, 2})
        self.assertTrue(all(c.metadata["source"] == "report.pdf" for c in chunks))
        self.assertTrue(all("chunk_id" in c.metadata for c in chunks))
        self.assertEqual(len({c.metadata["chunk_id"] for c in chunks}), len(chunks))
        self.assertFalse(any("First page" in c.page_content and "Second page" in c.page_content for c in chunks))

    def test_real_pdf_can_be_parsed(self):
        documents = parse_uploaded_file(Upload("sample.pdf", (ROOT / "test" / "test.pdf").read_bytes()))
        self.assertTrue(documents)
        self.assertEqual(documents[0].metadata["page"], 1)

    def test_docx_keeps_table_location_without_inventing_pages(self):
        word = WordDocument()
        word.add_paragraph("Project notes")
        word.add_table(rows=1, cols=1).cell(0, 0).text = "Revenue: 42"
        buffer = io.BytesIO()
        word.save(buffer)
        documents = parse_uploaded_file(Upload("notes.docx", buffer.getvalue()))
        self.assertEqual([d.metadata["location"] for d in documents], ["document text", "table 1"])
        self.assertTrue(all("page" not in d.metadata for d in documents))

    def test_csv_bom_and_quoted_fields(self):
        doc = parse_uploaded_file(Upload("sales.csv", b'\xef\xbb\xbfname,total\n"A, B",42\n'))[0]
        self.assertIn("A, B, 42", doc.page_content)
        self.assertEqual(doc.metadata["location"], "CSV rows")

    def test_bad_inputs_are_rejected(self):
        for upload in [Upload("empty.txt", b""), Upload("blank.txt", b" \n"),
                       Upload("bad.pdf", b"not a PDF"), Upload("bad.txt", b"\xff"),
                       Upload("unsupported.xlsx", b"data")]:
            with self.subTest(name=upload.name), self.assertRaises(Exception):
                parse_uploaded_file(upload)

    def test_scanned_pdf_has_actionable_failure(self):
        with patch("document_utils.PdfReader", return_value=SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda: None)])):
            with self.assertRaisesRegex(ValueError, "OCR"):
                parse_uploaded_file(Upload("scan.pdf", b"fake"))

    def test_partial_ingestion_reports_every_file_and_progress(self):
        events = []
        chunks, report = ingest_files([Upload("good.md", b"Useful notes."), Upload("bad.pdf", b"corrupt")],
                                     progress=lambda *args: events.append(args))
        self.assertTrue(chunks)
        self.assertEqual([row["status"] for row in report], ["Parsed", "Failed"])
        self.assertEqual([event[0] for event in events], [1, 2])
        self.assertEqual(report[0]["chunks"], len(chunks))
        self.assertTrue(report[1]["detail"])


class RetrievalTests(unittest.TestCase):
    def test_keyword_branch_recovers_identifier_missed_by_semantic_branch(self):
        chunks = get_chunks([
            Document(page_content="Warehouse identifier MARIGOLD-42", metadata={"source": "codes.txt"}),
            Document(page_content="Cats enjoy sleep.", metadata={"source": "cats.txt"}),
            Document(page_content="Birds fly in the sky.", metadata={"source": "birds.txt"}),
        ])
        vectorstore = SimpleNamespace(as_retriever=lambda **kwargs: StaticRetriever(documents=chunks[1:]))
        hybrid = get_hybrid_retriever(vectorstore, chunks)
        results = hybrid.invoke("MARIGOLD-42")
        self.assertIn(chunks[0], results)
        self.assertEqual(len({doc.metadata["chunk_id"] for doc in results}), len(results))

    def test_zero_keyword_matches_do_not_fabricate_retrieval(self):
        chunks = get_chunks("Cats enjoy sleep.")
        vectorstore = SimpleNamespace(as_retriever=lambda **kwargs: StaticRetriever(documents=[]))
        self.assertEqual(get_hybrid_retriever(vectorstore, chunks).invoke("unrelatedxyz"), [])

    def test_sources_are_attributed_and_deduplicated(self):
        chunk = get_chunks([Document(page_content="A verified excerpt.", metadata={"source": "real.pdf", "page": 2, "location": "page 2"})])[0]
        sources = get_sources([chunk, chunk])
        self.assertEqual(len(sources), 1)
        self.assertEqual(sources[0]["file"], "real.pdf")
        self.assertEqual(sources[0]["page"], 2)
        self.assertEqual(sources[0]["excerpt"], chunk.page_content)


class UITests(unittest.TestCase):
    def test_history_and_sources_survive_ordinary_rerun(self):
        app = AppTest.from_file(str(ROOT / "src" / "app.py")).run(timeout=30)
        app.session_state.chat_history = [
            {"role": "user", "content": "HISTORY_SENTINEL"},
            {"role": "assistant", "content": "ANSWER_SENTINEL", "sources": [
                {"file": "real.pdf", "location": "page 2", "excerpt": "EXCERPT_SENTINEL"}]},
        ]
        app.run(timeout=30)
        self.assertFalse(app.exception)
        self.assertEqual(len(app.chat_message), 2)
        self.assertIn("HISTORY_SENTINEL", [element.value for element in app.markdown])
        self.assertTrue(any("real.pdf" in element.value for element in app.text))
        app.toggle[0].set_value(True).run(timeout=30)
        self.assertEqual(len(app.chat_message), 2)
        next(button for button in app.button if button.label == "Clear Chat History").click().run(timeout=30)
        self.assertEqual(len(app.session_state.chat_history), 0)
        self.assertEqual(len(app.chat_message), 0)

    def test_empty_context_or_answer_never_uses_ungrounded_fallback(self):
        source = '''
import streamlit as st
from unittest.mock import patch
from langchain_core.documents import Document
from conversation import handle_question, NO_ANSWER
class Chain:
    def __init__(self, response): self.response = response
    def invoke(self, payload): return self.response
for response in ({'answer': 'invented', 'context': []},
                 {'answer': '', 'context': [Document(page_content='fact', metadata={'source':'real.txt'})]}):
    st.session_state.chat_history = []
    st.session_state.conversation = Chain(response)
    with patch('conversation.ChatOllama', side_effect=AssertionError('Unexpected model fallback')):
        assert handle_question('Original question', retrieval_question='Rewritten query')
    assert st.session_state.chat_history[0]['content'] == 'Original question'
    assert st.session_state.chat_history[-1]['content'] == NO_ANSWER
st.text('passed')
'''
        app = AppTest.from_string(source).run(timeout=30)
        self.assertFalse(app.exception, [e.value for e in app.exception])

    def test_model_failure_is_reported_without_a_fabricated_answer(self):
        source = '''
import streamlit as st
from conversation import handle_question
class Chain:
    def invoke(self, payload): raise ConnectionError('Ollama unavailable')
st.session_state.chat_history = []
st.session_state.conversation = Chain()
assert not handle_question('Question')
assert not st.session_state.chat_history
'''
        app = AppTest.from_string(source).run(timeout=30)
        self.assertFalse(app.exception)
        self.assertIn("Ollama unavailable", app.error[0].value)

    def test_embedding_failure_keeps_parsed_report_but_not_ready_state(self):
        source = '''
import io
import streamlit as st
from unittest.mock import patch
from document_utils import process_documents
class Upload(io.BytesIO): name = 'notes.txt'
with patch('document_utils.get_vectorstore', return_value=None):
    assert not process_documents([Upload(b'Useful notes')])
assert st.session_state.conversation is None
assert not st.session_state.processed
assert st.session_state.ingestion_report[0]['status'] == 'Parsed'
'''
        app = AppTest.from_string(source).run(timeout=30)
        self.assertFalse(app.exception, [e.value for e in app.exception])


if __name__ == "__main__":
    unittest.main()
