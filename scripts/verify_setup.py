"""Verify Streamlit startup, file readers, and optionally the live local RAG chain."""
import argparse
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from streamlit.testing.v1 import AppTest


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--models", action="store_true")
    args = parser.parse_args()
    app = AppTest.from_file(str(ROOT / "src" / "app.py")).run(timeout=30)
    assert not app.exception, [exception.value for exception in app.exception]
    assert app.header[0].value == "Chat with Your Documents :books:"
    print("PASS: Streamlit UI starts without errors.")

    from document_utils import get_text_from_file
    from docx import Document

    class Upload(io.BytesIO):
        def __init__(self, content, mime_type):
            super().__init__(content)
            self.type = mime_type

    assert "Cats" in get_text_from_file(Upload((ROOT / "test" / "test.txt").read_bytes(), "text/plain"))
    assert get_text_from_file(Upload((ROOT / "test" / "test.pdf").read_bytes(), "application/pdf"))
    assert "blue" in get_text_from_file(Upload(b"name,color\nsky,blue\n", "text/csv"))
    word = Document()
    word.add_paragraph("A document paragraph.")
    word.add_table(rows=1, cols=1).cell(0, 0).text = "A table value."
    buffer = io.BytesIO()
    word.save(buffer)
    text = get_text_from_file(Upload(buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
    assert "document paragraph" in text and "table value" in text
    print("PASS: PDF, TXT, CSV and DOCX readers extract text.")

    if args.models:
        source = '''
import io
import streamlit as st
from document_utils import process_documents
from conversation import handle_question
from prompt_refiner import refine_prompt_with_llm

class Upload(io.BytesIO):
    type = "text/plain"
    name = "warehouse.txt"

st.session_state.chat_history = []
sample = Upload(b"The project's sample warehouse is in Jaipur. Its access code is MARIGOLD-42.")
if not process_documents([sample]):
    raise RuntimeError("Document indexing failed")
handle_question("What is the warehouse access code in the document?")
if not st.session_state.chat_history:
    raise RuntimeError("No answer returned")
answer = st.session_state.chat_history[-1]["content"]
assert "MARIGOLD-42" in answer, answer
sources = st.session_state.chat_history[-1]['sources']
assert sources and sources[0]['file'] == 'warehouse.txt', sources
assert 'MARIGOLD-42' in sources[0]['excerpt']
st.text(answer)
handle_question("Which city is the warehouse in?")
assert "Jaipur" in st.session_state.chat_history[-1]['content']
handle_question("What is the warehouse manager's favorite ice cream flavor?")
missing = st.session_state.chat_history[-1]['content'].lower()
assert any(term in missing for term in ("don't know", "does not", "do not", "not provided", "no information", "not mentioned", "not contain")), missing
refined = refine_prompt_with_llm("What is the warehouse access code?")
assert refined.strip()
st.text(refined)
'''
        rag = AppTest.from_string(source).run(timeout=300)
        assert not rag.exception, [exception.value for exception in rag.exception]
        assert not rag.error, [error.value for error in rag.error]
        print("PASS: Local embeddings, hybrid retrieval, attributed excerpts, follow-up chat and missing-answer behavior.")
        print("Answer:", rag.text[0].value)
        print("PASS: Local prompt refinement returns text.")


if __name__ == "__main__":
    main()
