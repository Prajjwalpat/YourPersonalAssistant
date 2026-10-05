"""Extract independent, attributable documents and report per-file ingestion."""
import csv
import hashlib
import io
from pathlib import Path
from typing import Callable, List, Optional

import streamlit as st
from docx import Document as WordDocument
from langchain_core.documents import Document
from PyPDF2 import PdfReader

from conversation import get_conversationchain
from text_processing import get_chunks, get_hybrid_retriever, get_vectorstore

TEXT_EXTENSIONS = {
    "txt", "md", "markdown", "py", "js", "ts", "jsx", "tsx", "java", "go",
    "rs", "c", "cpp", "h", "hpp", "cs", "sql", "sh", "ps1", "r", "rb",
    "json", "yaml", "yml", "toml", "html", "css",
}
SUPPORTED_EXTENSIONS = ["pdf", "docx", "csv", *sorted(TEXT_EXTENSIONS)]


def parse_uploaded_file(uploaded_file) -> List[Document]:
    """Read fresh bytes on every call; PDF page numbers are one-based."""
    name = getattr(uploaded_file, "name", "")
    extension = Path(name).suffix.lower().lstrip(".")
    if not extension:
        extension = {
            "application/pdf": "pdf", "text/plain": "txt", "text/csv": "csv",
            "text/markdown": "md", "application/vnd.ms-excel": "csv",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        }.get(getattr(uploaded_file, "type", ""), "")
    name = Path(name).name if name else f"Uploaded file.{extension or 'unknown'}"
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file extension: .{extension or 'unknown'}")
    if hasattr(uploaded_file, "getvalue"):
        content = uploaded_file.getvalue()
    else:
        uploaded_file.seek(0)
        content = uploaded_file.read()
    if not content:
        raise ValueError("The file is empty.")
    file_id = hashlib.sha256(name.encode() + content).hexdigest()[:20]

    def unit(text: str, location: str, **metadata) -> Document:
        return Document(page_content=text, metadata={
            "source": name, "file_id": file_id, "location": location,
            "format": extension, **metadata,
        })

    if extension == "pdf":
        reader = PdfReader(io.BytesIO(content))
        documents = [unit(page.extract_text() or "", f"page {number}", page=number)
                     for number, page in enumerate(reader.pages, 1)]
    elif extension == "docx":
        word = WordDocument(io.BytesIO(content))
        documents = [unit("\n".join(paragraph.text for paragraph in word.paragraphs), "document text")]
        for number, table in enumerate(word.tables, 1):
            text = "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            documents.append(unit(text, f"table {number}", table=number))
    elif extension == "csv":
        text = content.decode("utf-8-sig")
        rows = csv.reader(io.StringIO(text))
        documents = [unit("\n".join(", ".join(row) for row in rows), "CSV rows")]
    else:
        documents = [unit(content.decode("utf-8-sig"), "file text")]
    documents = [doc for doc in documents if doc.page_content.strip()]
    if not documents:
        raise ValueError("No readable text found. Scanned/image-only PDFs need OCR before upload.")
    return documents


def get_text_from_file(uploaded_file) -> Optional[str]:
    """Compatibility wrapper used by the original smoke-check script."""
    try:
        return "\n".join(doc.page_content for doc in parse_uploaded_file(uploaded_file))
    except Exception as error:
        st.error(f"Cannot read the file: {error}")
        return None


def ingest_files(files, progress: Optional[Callable] = None):
    chunks, report = [], []
    for index, uploaded_file in enumerate(files, 1):
        name = getattr(uploaded_file, "name", f"File {index}")
        try:
            documents = parse_uploaded_file(uploaded_file)
            file_chunks = get_chunks(documents)
            chunks.extend(file_chunks)
            report.append({"file": name, "status": "Parsed", "units": len(documents),
                           "chunks": len(file_chunks), "detail": "Text extracted successfully."})
        except Exception as error:
            report.append({"file": name, "status": "Failed", "units": 0, "chunks": 0, "detail": str(error)})
        if progress:
            progress(index, len(files), name)
    return chunks, report


def process_documents(docs, progress: Optional[Callable] = None) -> bool:
    st.session_state.conversation = None
    st.session_state.processed = False
    chunks, report = ingest_files(docs, progress)
    st.session_state.ingestion_report = report
    st.session_state.chunk_count = len(chunks)
    if not chunks:
        st.error("No readable documents were indexed. Check the per-file report.")
        return False
    try:
        vectorstore = get_vectorstore(chunks)
        if vectorstore is None:
            return False
        retriever = get_hybrid_retriever(vectorstore, chunks)
        chain = get_conversationchain(vectorstore, retriever=retriever)
        if chain is None:
            return False
        st.session_state.conversation = chain
        for row in report:
            if row["status"] == "Parsed":
                row["status"] = "Ready"
                row["detail"] = "Indexed for semantic and keyword search."
        return True
    except Exception as error:
        st.error(f"Document indexing failed: {error}")
        return False
